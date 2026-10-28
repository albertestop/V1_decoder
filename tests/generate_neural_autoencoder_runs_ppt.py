#!/usr/bin/env python3
"""Create a PowerPoint table comparing neural-autoencoder runs.

Examples
--------
python tests/generate_neural_autoencoder_runs_ppt.py 203 204 205 \\
    --title "TAE_v1_01_1 - compression study"

# Supply presentation labels in the same order as the run numbers.
python tests/generate_neural_autoencoder_runs_ppt.py 203 204 \\
    --labels "1024 -> 512, t_w = 10, id_w = 0.1" "256, t_w = 1" \\
    --output outputs/neural_autoencoder/comparison.pptx

Requires ``python-pptx`` in the Python environment used to run the script:
    python -m pip install python-pptx
"""

from __future__ import annotations

import argparse
import importlib.util
import sys
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RUNS_DIRECTORY = PROJECT_ROOT / "outputs" / "neural_autoencoder"
DEFAULT_OUTPUT = RUNS_DIRECTORY / "run_comparison.pptx"
MISSING_VALUE = "—"


def load_summary_module():
    """Load the existing summary script without requiring ``tests`` to be a package."""
    path = Path(__file__).with_name("summarize_neural_autoencoder_run.py")
    spec = importlib.util.spec_from_file_location("run_summary", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not load run summary helper from {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def load_toml(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    try:
        import tomllib
    except ModuleNotFoundError:  # Python 3.10 and older
        import tomli as tomllib
    with path.open("rb") as file:
        return tomllib.load(file)


def format_metric(error: str, std: str) -> str:
    if error == "non-existent":
        return MISSING_VALUE
    return f"{error} ± {std}" if std != "non-existent" else error


def default_label(run_identifier: str) -> str:
    """Use only the run number unless an explicit --labels value is supplied."""
    return run_identifier


def collect_row(run_identifier: str, label: str | None) -> list[str]:
    """Collect the exact metrics printed by summarize_neural_autoencoder_run.py."""
    summary = load_summary_module()
    run_dir = RUNS_DIRECTORY / f"run_{run_identifier}"
    validation_path = run_dir / "validation_error_stats.json"
    counts_path = run_dir / "high_mae_trial_plots" / "plot_counts.txt"
    history_path = run_dir / "history.json"

    if validation_path.is_file():
        stats = summary.load_json(validation_path)
        token_1, token_2 = stats["token_1"], stats["token_2"]
        token_1_value = format_metric(
            summary.format_value(token_1["pred_global_distance_to_target"]),
            summary.format_value(token_1["global_distance_std"]),
        )
        token_2_value = format_metric(
            summary.format_value(token_2["pred_global_distance_to_target"]),
            summary.format_value(token_2["global_distance_std"]),
        )
        misalignment = summary.format_value(token_1["mean_align_dist"])
    else:
        token_1_value = token_2_value = misalignment = MISSING_VALUE

    if counts_path.is_file():
        counts = summary.read_outlier_counts(counts_path)
        outliers = f"{counts.get('token_1', 'not found')} | {counts.get('token_2', 'not found')}"
    else:
        outliers = MISSING_VALUE

    if history_path.is_file():
        history = summary.load_json(history_path)
        if history:
            best = min(history, key=lambda epoch: epoch["val_loss"])
            early_stop = f"{best['epoch']:g}"
            overfitting = "x" + summary.format_value(best["val_loss"] / best["train_loss"])
            time_overfitting = summary.format_value(best["val_loss_time"] / best["train_loss_time"])
        else:
            early_stop = overfitting = time_overfitting = MISSING_VALUE
    else:
        early_stop = overfitting = time_overfitting = MISSING_VALUE

    return [
        label or default_label(run_identifier), token_1_value, token_2_value, outliers,
        early_stop, misalignment, overfitting, time_overfitting,
    ]


HEADERS = [
    "Run / configuration", "Token 1", "Token 2", "Outliers\n(token 1 | token 2)",
    "Early\nstop", "Token 1\nmisalignment", "Overfitting\n(val/train)",
    "Time overfitting\n(val/train)",
]


def set_cell_border(cell, color: str = "A6A6A6") -> None:
    """Add thin grey borders, matching the reference table's visual weight."""
    from pptx.oxml.xmlchemy import OxmlElement

    tc_pr = cell._tc.get_or_add_tcPr()
    for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        line = OxmlElement(edge)
        line.set("w", "12700")
        solid = OxmlElement("a:solidFill")
        rgb = OxmlElement("a:srgbClr")
        rgb.set("val", color)
        solid.append(rgb)
        line.append(solid)
        tc_pr.append(line)


def add_table_slide(presentation, title: str, rows: list[list[str]], slide_number: int, total_slides: int) -> None:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.32), Inches(0.20), Inches(12.70), Inches(0.48))
    title_frame = title_box.text_frame
    title_frame.clear()
    paragraph = title_frame.paragraphs[0]
    paragraph.text = title if total_slides == 1 else f"{title} ({slide_number}/{total_slides})"
    paragraph.font.name = "Arial"
    paragraph.font.size = Pt(28)
    paragraph.font.bold = False
    paragraph.font.color.rgb = RGBColor(0, 0, 0)

    row_count = len(rows) + 1
    table_shape = slide.shapes.add_table(
        row_count, len(HEADERS), Inches(0.22), Inches(0.88), Inches(12.88), Inches(6.34)
    )
    table = table_shape.table
    widths = [3.25, 2.22, 2.02, 1.54, 0.62, 1.38, 1.08, 0.77]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)

    for row_index in range(row_count):
        table.rows[row_index].height = Inches(0.42 if row_index == 0 else min(0.56, 6.34 / row_count))
    for column_index, header in enumerate(HEADERS):
        cell = table.cell(0, column_index)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(242, 242, 242)
        set_cell_border(cell)

    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
            set_cell_border(cell)

    for row in table.rows:
        for cell in row.cells:
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.03)
            cell.margin_right = Inches(0.02)
            cell.margin_top = cell.margin_bottom = Inches(0.01)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                paragraph.font.name = "Arial"
                paragraph.font.size = Pt(10 if row == table.rows[0] else 9)
                paragraph.font.color.rgb = RGBColor(0, 0, 0)


def create_presentation(rows: list[list[str]], title: str, output_path: Path) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ModuleNotFoundError as error:
        raise SystemExit(
            "python-pptx is required. Install it in this Python environment with: "
            f"{sys.executable} -m pip install python-pptx"
        ) from error

    # At this font size, 12 data rows fit comfortably on a 16:9 slide.
    rows_per_slide = 12
    chunks = [rows[index:index + rows_per_slide] for index in range(0, len(rows), rows_per_slide)]
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    for index, chunk in enumerate(chunks, start=1):
        add_table_slide(presentation, title, chunk, index, len(chunks))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("runs", nargs="+", help="Run numbers or names (for example: 203 run_204).")
    parser.add_argument("--labels", nargs="+", help="Optional presentation labels, one per run, in the same order.")
    parser.add_argument("--title", default="Neural autoencoder run comparison", help="Slide title.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help=f"Output .pptx path (default: {DEFAULT_OUTPUT}).")
    arguments = parser.parse_args()
    if arguments.labels is not None and len(arguments.labels) != len(arguments.runs):
        parser.error("--labels must contain exactly one label for each run.")

    run_ids = [run.removeprefix("run_") for run in arguments.runs]
    labels = arguments.labels or [None] * len(run_ids)
    rows = [collect_row(run_id, label) for run_id, label in zip(run_ids, labels)]
    create_presentation(rows, arguments.title, arguments.output)
    print(f"Created {arguments.output.resolve()} with {len(rows)} run(s).")


if __name__ == "__main__":
    main()
